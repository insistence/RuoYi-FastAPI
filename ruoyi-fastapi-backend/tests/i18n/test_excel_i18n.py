import io
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pandas as pd
import pytest
from fastapi import UploadFile
from openpyxl import load_workbook

from common.context import LocaleContext, RequestContext
from exceptions.exception import ServiceException
from i18n.message import I18nMessage
from module_admin.dao.user_dao import UserDao
from module_admin.service.config_service import ConfigService
from module_admin.service.dict_service import DictTypeService
from module_admin.service.user_service import UserService
from module_identity.service.identity_service import IdentitySubjectService
from utils.excel_util import ExcelUtil
from utils.user_import_util import UserImportUtil


@pytest.mark.asyncio
@pytest.mark.parametrize('locale', ['zh-CN', 'en-US'])
async def test_published_user_templates_round_trip_independently_of_interface_language(
    locale: str, monkeypatch: pytest.MonkeyPatch
) -> None:
    token = RequestContext.set_current_locale(LocaleContext(locale, 'v1'))
    try:
        workbook = load_workbook(io.BytesIO(await UserService.get_user_import_template_services()))
    finally:
        RequestContext.reset_current_locale(token)
    headers = [cell.value for cell in workbook.active[1]]
    assert UserImportUtil.normalize_headers(headers) == list(UserImportUtil.HEADER_ALIASES)
    validations = list(workbook.active.data_validations.dataValidation)
    options = [validation.formula1.strip('"').split(',') for validation in validations]
    assert [UserImportUtil.normalize_value('sex', label) for label in options[0]] == ['0', '1', '2']
    assert [UserImportUtil.normalize_value('status', label) for label in options[1]] == ['0', '1']
    workbook.active.append([100, 'ada', '自定义昵称', 'ada@example.com', '01234567890', options[0][0], options[1][0]])
    stream = io.BytesIO()
    workbook.save(stream)
    sheet = pd.read_excel(io.BytesIO(stream.getvalue()), header=None, dtype=object)
    sheet.columns = UserImportUtil.normalize_headers(sheet.iloc[0].tolist())
    row = sheet.iloc[1].copy()
    UserService._set_row_sex_value(row)
    UserService._set_row_status_value(row)
    assert row['sex'] == '0'
    assert row['status'] == '0'
    assert row['nick_name'] == '自定义昵称'
    assert row['phonenumber'] == '01234567890'
    monkeypatch.setattr(ConfigService, 'query_config_list_from_cache_services', AsyncMock(return_value='ValidPass123'))
    monkeypatch.setattr(UserService, 'validate_password_services', AsyncMock())
    monkeypatch.setattr(UserDao, 'get_user_by_info', AsyncMock(return_value=None))
    add_user = AsyncMock(return_value=SimpleNamespace(user_id=77))
    monkeypatch.setattr(UserDao, 'add_user_dao', add_user)
    monkeypatch.setattr(IdentitySubjectService, 'create_for_new_user', AsyncMock())
    db = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock())
    result = await UserService.batch_import_user_services(
        SimpleNamespace(app=SimpleNamespace(state=SimpleNamespace(redis=object()))),
        db,
        UploadFile(file=io.BytesIO(stream.getvalue()), filename='users.xlsx'),
        False,
        SimpleNamespace(user=SimpleNamespace(admin=True, user_name='admin')),
        None,
        None,
    )
    created = add_user.call_args.args[1]
    assert created.user_name == 'ada'
    assert created.nick_name == '自定义昵称'
    assert created.sex == '0'
    assert created.status == '0'
    assert result.msg_key == 'common.imported'
    db.commit.assert_awaited_once()


def test_import_rejects_duplicate_aliases_missing_columns_and_unknown_enums() -> None:
    headers = [aliases[0] for aliases in UserImportUtil.HEADER_ALIASES.values()]
    with pytest.raises(ServiceException) as duplicate:
        UserImportUtil.normalize_headers([*headers, 'Login name'])
    assert duplicate.value.msg_key == 'import.duplicateColumn'
    with pytest.raises(ServiceException) as missing:
        UserImportUtil.normalize_headers(headers[:-1])
    assert missing.value.msg_key == 'import.missingColumns'
    with pytest.raises(ServiceException) as invalid:
        UserImportUtil.normalize_value('status', '未发布状态')
    assert invalid.value.msg_key == 'import.invalidValue'


def test_export_uses_snapshot_and_keeps_business_content_and_legacy_headers() -> None:
    mapping = {'value': I18nMessage('用户名称', msg_key='export.userName')}
    for context, expected in [
        (LocaleContext('en-US', 'v1'), 'Username'),
        (LocaleContext('en-US', 'legacy'), '用户名称'),
    ]:
        binary = ExcelUtil.export_list2excel([{'value': '用户提供的内容'}], mapping, locale_context=context)
        sheet = load_workbook(io.BytesIO(binary)).active
        assert sheet.cell(1, 1).value == expected
        assert sheet.cell(2, 1).value == '用户提供的内容'


@pytest.mark.asyncio
async def test_dictionary_export_localizes_name_and_status_without_changing_type() -> None:
    token = RequestContext.set_current_locale(LocaleContext('en-US', 'v1'))
    try:
        data = await DictTypeService.export_dict_type_list_services(
            [
                {
                    'dictId': 1,
                    'dictName': '用户性别',
                    'dictNameI18n': {'en-US': 'User gender'},
                    'dictType': 'sys_user_sex',
                    'status': '0',
                },
            ]
        )
    finally:
        RequestContext.reset_current_locale(token)
    sheet = load_workbook(io.BytesIO(data)).active
    row = dict(zip([cell.value for cell in sheet[1]], [cell.value for cell in sheet[2]], strict=True))
    assert row['Dictionary name'] == 'User gender'
    assert row['Dictionary type'] == 'sys_user_sex'
    assert row['Status'] == 'Active'
